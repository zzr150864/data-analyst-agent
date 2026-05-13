import os
from datetime import datetime
from hello_agents import BaseTool
from session import session
from utils.file_utils import ensure_output_dir, output_path


class GenerateMarkdownReportTool(BaseTool):
    name = "generate_markdown_report"
    description = "Generate a Markdown analysis report. Provide title, data_summary, analysis_findings, visualization_descriptions, and optional conclusions. Saves to outputs/."

    def run(
        self,
        title: str,
        data_summary: str,
        analysis_findings: str,
        visualization_descriptions: str,
        conclusions: str = "",
    ) -> str:
        ensure_output_dir()
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = output_path(f"report_{timestamp}.md")

        chart_links = []
        for p in session.chart_paths:
            fname = p.split("/")[-1].split("\\")[-1]
            if p.endswith(".png"):
                chart_links.append(f"![{fname}]({fname})")
            elif p.endswith(".html"):
                chart_links.append(f"- [Interactive: {fname}]({fname})")

        content = f"""# {title}

**Generated:** {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
**Source File:** {session.file_path}

---

## 1. Data Overview

{data_summary}

---

## 2. Statistical Analysis

{analysis_findings}

---

## 3. Visualizations

{visualization_descriptions}

{chr(10).join(chart_links)}

---

## 4. Conclusions

{conclusions or 'See analysis findings above for key takeaways.'}
"""
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        session.report_paths.append(path)
        return f"Markdown report saved: {path}"


class GenerateHTMLReportTool(BaseTool):
    name = "generate_html_report"
    description = "Generate a styled HTML analysis report. Provide title, data_summary, analysis_findings, visualization_descriptions, and optional conclusions. Saves to outputs/."

    def run(
        self,
        title: str,
        data_summary: str,
        analysis_findings: str,
        visualization_descriptions: str,
        conclusions: str = "",
    ) -> str:
        ensure_output_dir()
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = output_path(f"report_{timestamp}.html")

        chart_html = []
        for p in session.chart_paths:
            fname = p.split("/")[-1].split("\\")[-1]
            if p.endswith(".png"):
                chart_html.append(f'<img src="{fname}" alt="{fname}" />')
            elif p.endswith(".html"):
                chart_html.append(
                    f'<p><a href="{fname}" target="_blank">Interactive Chart: {fname}</a></p>'
                )

        html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>{title}</title>
<style>
  body {{
    font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', sans-serif;
    max-width: 960px;
    margin: 0 auto;
    padding: 2rem;
    color: #333;
    line-height: 1.6;
  }}
  h1 {{ border-bottom: 3px solid #4A90D9; padding-bottom: 0.5rem; color: #2c3e50; }}
  h2 {{ color: #4A90D9; margin-top: 2.5rem; }}
  pre {{
    background: #f5f7fa;
    padding: 1.2rem;
    border-radius: 8px;
    overflow-x: auto;
    font-size: 0.9rem;
    line-height: 1.5;
  }}
  .meta {{ color: #888; font-size: 0.9rem; margin-bottom: 2rem; }}
  img {{ max-width: 100%; border: 1px solid #e0e0e0; border-radius: 6px; margin: 1rem 0; box-shadow: 0 2px 8px rgba(0,0,0,0.1); }}
  a {{ color: #4A90D9; }}
</style>
</head>
<body>
<h1>{title}</h1>
<p class="meta">Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Source: {session.file_path}</p>

<h2>1. Data Overview</h2>
<pre>{data_summary}</pre>

<h2>2. Statistical Analysis</h2>
<pre>{analysis_findings}</pre>

<h2>3. Visualizations</h2>
<p>{visualization_descriptions}</p>
{chr(10).join(chart_html)}

<h2>4. Conclusions</h2>
<p>{conclusions or 'See analysis findings above for key takeaways.'}</p>
</body>
</html>"""
        with open(path, "w", encoding="utf-8") as f:
            f.write(html)
        session.report_paths.append(path)
        return f"HTML report saved: {path}"


class GeneratePDFReportTool(BaseTool):
    name = "generate_pdf_report"
    description = "Generate a PDF analysis report with embedded charts. Provide title, data_summary, analysis_findings, visualization_descriptions, and optional conclusions. Saves to outputs/."

    def run(
        self,
        title: str,
        data_summary: str,
        analysis_findings: str,
        visualization_descriptions: str,
        conclusions: str = "",
    ) -> str:
        ensure_output_dir()
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = output_path(f"report_{timestamp}.pdf")

        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()

            # Title
            h = doc.add_heading(title, level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER

            doc.add_paragraph(
                f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | "
                f"Source: {session.file_path}"
            ).style = doc.styles["Normal"]

            # Sections
            sections = [
                ("1. Data Overview", data_summary),
                ("2. Statistical Analysis", analysis_findings),
                ("3. Visualizations", visualization_descriptions),
                ("4. Conclusions", conclusions or "See analysis findings above."),
            ]

            for section_title, content in sections:
                doc.add_heading(section_title, level=1)
                for paragraph in content.split("\n"):
                    paragraph = paragraph.strip()
                    if paragraph:
                        doc.add_paragraph(paragraph)

            # Embed PNG charts
            png_charts = [p for p in session.chart_paths if p.lower().endswith(".png")]
            if png_charts:
                doc.add_heading("Charts", level=1)
                for chart_path in png_charts:
                    try:
                        fname = chart_path.split("\\")[-1].split("/")[-1]
                        doc.add_paragraph(fname, style="Caption")
                        doc.add_picture(chart_path, width=Inches(5.5))
                        doc.add_paragraph()  # spacing
                    except Exception as e:
                        doc.add_paragraph(f"(Chart: {chart_path} — {e})")

            # Save .docx first
            docx_path = path.replace(".pdf", ".docx")
            doc.save(docx_path)

            # Convert to PDF
            from docx2pdf import convert
            convert(docx_path, path)

            # Clean up .docx
            try:
                os.remove(docx_path)
            except Exception:
                pass

            session.report_paths.append(path)
            return f"PDF report saved: {path}"

        except Exception as e:
            # Fallback: simple PDF via matplotlib PdfPages
            try:
                from matplotlib.backends.backend_pdf import PdfPages
                import matplotlib.pyplot as plt

                with PdfPages(path) as pdf:
                    # Title page
                    fig, ax = plt.subplots(figsize=(8.5, 11))
                    ax.axis("off")
                    ax.text(0.5, 0.95, title, transform=ax.transAxes, fontsize=20,
                            ha="center", va="top", fontweight="bold")
                    ax.text(0.5, 0.90, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
                            transform=ax.transAxes, fontsize=10, ha="center", va="top",
                            color="gray")
                    ax.text(0.05, 0.85, data_summary[:3000], transform=ax.transAxes,
                            fontsize=8, va="top", family="monospace")
                    ax.text(0.05, 0.55, analysis_findings[:3000], transform=ax.transAxes,
                            fontsize=8, va="top", family="monospace")
                    ax.text(0.05, 0.25, conclusions or "", transform=ax.transAxes,
                            fontsize=10, va="top")
                    pdf.savefig(fig)
                    plt.close(fig)

                    # Chart pages
                    for chart_path in session.chart_paths:
                        if chart_path.lower().endswith(".png"):
                            try:
                                from PIL import Image
                                img = Image.open(chart_path)
                                fig, ax = plt.subplots(figsize=(8.5, 11))
                                ax.axis("off")
                                ax.imshow(img)
                                ax.set_title(chart_path.split("\\")[-1].split("/")[-1],
                                            fontsize=10)
                                pdf.savefig(fig)
                                plt.close(fig)
                            except Exception:
                                pass

                session.report_paths.append(path)
                return f"PDF report saved (matplotlib fallback): {path}"
            except Exception as fallback_err:
                return f"PDF generation failed: {e}. Fallback error: {fallback_err}"
